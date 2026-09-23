"""Artículo IEEE editable con resultados verificados e imágenes del experimento."""
import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/ieee"
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.75), Inches(1)
sec.left_margin = sec.right_margin = Inches(.625)
sec.header_distance = sec.footer_distance = Inches(.3)
styles = doc.styles
for name in ["Normal", "Title", "Heading 1", "Heading 2", "Caption"]:
    style = styles[name]
    style.font.name = "Times New Roman"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in list(fonts.attrib):
        if "Theme" in attr:
            del fonts.attrib[attr]
    for attr in ["ascii", "hAnsi", "eastAsia", "cs"]:
        fonts.set(qn("w:" + attr), "Times New Roman")
    for borders in list(style.element.iter(qn("w:pBdr"))):
        borders.getparent().remove(borders)
normal = styles["Normal"]
normal.font.size = Pt(10)
normal.paragraph_format.line_spacing = 1
normal.paragraph_format.first_line_indent = Inches(.14)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.widow_control = True
for name in ["Heading 1", "Heading 2"]:
    styles[name].font.size = Pt(10)
    styles[name].font.bold = False
    styles[name].paragraph_format.space_before = Pt(9)
    styles[name].paragraph_format.space_after = Pt(5)
    styles[name].paragraph_format.first_line_indent = Inches(0)
    styles[name].paragraph_format.keep_with_next = True
styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
styles["Heading 2"].font.italic = True
styles["Caption"].font.size = Pt(8)
styles["Caption"].font.italic = False
styles["Caption"].font.bold = False
styles["Caption"].paragraph_format.first_line_indent = Inches(0)
styles["Caption"].paragraph_format.space_after = Pt(7)
styles["Caption"].paragraph_format.keep_with_next = False


def columns(n, new_page=False):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE if new_page else WD_SECTION_START.CONTINUOUS)
    cols = section._sectPr.find(qn("w:cols"))
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), "360")
    return section


def p(text, bold=False):
    text = re.sub(r"\[([1-5])\]", lambda m: "[" + {"1":"1", "2":"3", "3":"4", "4":"2", "5":"5"}[m.group(1)] + "]", text)
    paragraph = doc.add_paragraph(text)
    if bold:
        paragraph.runs[0].bold = True
        paragraph.paragraph_format.first_line_indent = Inches(0)
        paragraph.paragraph_format.space_after = Pt(5)
        for run in paragraph.runs:
            run.font.size = Pt(9)
    return paragraph


def column_break():
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    paragraph.add_run().add_break(WD_BREAK.COLUMN)


def heading(text):
    doc.add_paragraph(text, "Heading 1")


def figure(path, caption, width=7.15):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    pic = run.add_picture(str(ROOT / path), width=Inches(width))
    pic._inline.docPr.set("descr", caption)
    doc.add_paragraph(caption, "Caption")


title = doc.add_paragraph("Clasificación de la actividad de moléculas frente a MMP-9 mediante aprendizaje automático", "Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(13)
for run in title.runs:
    run.font.size = Pt(24)
    run.font.bold = False
author = doc.add_paragraph("Geronimo Ortiz Porras   y   Samuel Tomas Rubiano Meneses")
author.alignment = WD_ALIGN_PARAGRAPH.CENTER
author.paragraph_format.first_line_indent = Inches(0)
author.runs[0].font.size = Pt(11)
affiliation = doc.add_paragraph("Universidad del Rosario")
affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
affiliation.paragraph_format.first_line_indent = Inches(0)
affiliation.paragraph_format.space_after = Pt(12)
affiliation.runs[0].font.size = Pt(10)
columns(2)
p("Resumen—Este trabajo evaluó la capacidad de tres modelos de aprendizaje automático para distinguir moléculas activas e inactivas frente a la metaloproteinasa de matriz 9 humana, MMP-9. A partir de ChEMBL se construyó un conjunto de 913 moléculas, representadas mediante cadenas SMILES vectorizadas y descriptores PaDEL. Se compararon regresión logística, Random Forest y XGBoost en 100 repeticiones por combinación, para un total de 600 evaluaciones. XGBoost con PaDEL alcanzó una accuracy de 93,46%, un F1-score de 95,88%, una sensibilidad de 96,67% y una especificidad de 81,62%. Random Forest con PaDEL obtuvo el mayor F1 medio, mientras que la regresión logística con PaDEL reconoció mejor las moléculas inactivas. Los resultados muestran utilidad para una selección computacional preliminar, aunque su alcance está limitado por la distribución de los datos y la ausencia de validación con familias químicas independientes.", True)
p("Palabras clave—MMP-9, bioactividad, aprendizaje automático, XGBoost, SMILES, PaDEL.", True)
heading("I. INTRODUCCIÓN")
p("La metaloproteinasa de matriz 9, conocida como MMP-9 o gelatinasa B, es una proteína humana asociada con la degradación de componentes de la matriz extracelular. Su registro en UniProt corresponde a P14780 [1]. En este trabajo se estudió la relación entre la estructura de moléculas evaluadas frente a esta proteína y la actividad inhibitoria registrada en una base de datos pública.")
p("La información disponible permite explorar si un modelo puede reconocer patrones compartidos por moléculas activas. Esta aproximación puede apoyar la priorización de compuestos que merecen una revisión posterior. Sin embargo, una clasificación favorable no establece por sí sola seguridad, selectividad ni eficacia terapéutica. El objetivo del ejercicio fue comparar alternativas de representación molecular y de clasificación bajo un mismo procedimiento de evaluación.")
p("Se dio especial atención a XGBoost, un método que combina árboles de decisión construidos de forma sucesiva [4]. Para interpretar su rendimiento se incluyeron Random Forest y regresión logística como referencias. Además de la proporción total de aciertos, se midió la capacidad de encontrar moléculas activas y de reconocer las inactivas, porque ambos tipos de error tienen consecuencias distintas al seleccionar candidatos.")
heading("II. METODOLOGÍA")
doc.add_paragraph("A. Selección y preparación de los datos", "Heading 2")
p("Se consultó ChEMBL, versión 37, el 22 de septiembre de 2026, usando el objetivo CHEMBL321, identificado como MMP-9 de Homo sapiens [2]. La descarga inicial reunió 3.967 registros de IC50. Esta medida expresa la concentración necesaria para reducir a la mitad la actividad observada en un ensayo; dentro de condiciones comparables, un valor menor indica mayor potencia inhibitoria.")
p("Se conservaron valores exactos, positivos y expresados en nanomolar, procedentes de ensayos de unión con asignación directa a la proteína y nivel de confianza 9. Se descartaron registros con alertas de validez, posibles duplicados o condiciones que no cumplían estos criterios. Quedaron 1.164 mediciones elegibles. Las estructuras se normalizaron para agrupar representaciones equivalentes de una misma molécula y se utilizó la mediana cuando existían varias mediciones.")
p("Una molécula con registros tanto activos como inactivos fue excluida por contradicción. El conjunto curado quedó formado por 1.113 moléculas. Se definió como activa una molécula con IC50 menor o igual a 1.000 nM y como inactiva una con IC50 mayor o igual a 10.000 nM. Las 200 moléculas situadas entre ambos límites se conservaron en los archivos, pero no entraron en la clasificación. Así, el análisis se realizó sobre 719 activas y 194 inactivas. Estos umbrales delimitan el problema estudiado y no constituyen una clasificación clínica.")
doc.add_paragraph("B. Representación y evaluación", "Heading 2")
p("Cada molécula se describió de dos maneras. La primera convirtió fragmentos de dos a cuatro caracteres de su cadena SMILES en valores numéricos, mediante una ponderación TF-IDF y un máximo de 2.048 características. La segunda utilizó 1.444 descriptores PaDEL de una y dos dimensiones, que resumen propiedades y rasgos de la estructura molecular [3]. Ambas representaciones correspondieron exactamente a las mismas moléculas.")
p("En cada repetición se reservaron 183 moléculas para prueba y se emplearon 730 para entrenamiento, manteniendo la proporción de clases. Dentro del entrenamiento se hizo otra separación para elegir entre dos configuraciones de cada modelo según el F1 de validación. Después se reajustó el modelo con las 730 moléculas y se evaluó sobre las 183 reservadas. Las configuraciones se describen en la sección III; las mismas particiones se usaron para las seis combinaciones.")

columns(1, new_page=True)
heading("III. DISEÑO EXPERIMENTAL Y RESULTADOS")
figure("reports/imagenes/metricas_todas.png", "Fig. 1. Rendimiento de los tres modelos con SMILES y PaDEL. Las barras representan la media de 100 repeticiones y las líneas negras, la desviación estándar muestral. La clase positiva es la molécula activa.", 7.15)
cap = doc.add_paragraph("TABLA I\nRENDIMIENTO MEDIO Y DESVIACIÓN ESTÁNDAR EN PORCENTAJE", "Caption")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.keep_with_next = True
table = doc.add_table(rows=1, cols=6)
table.autofit = False
widths = [1.0, 1.43, 1.19, 1.19, 1.19, 1.25]
headers = ["Representación", "Modelo", "Accuracy", "F1-score", "Sensibilidad", "Especificidad"]
for cell, width, text in zip(table.rows[0].cells, widths, headers):
    cell.width = Inches(width)
    cell.text = text
with (ROOT / "results/main/summary.csv").open(encoding="utf-8", newline="") as stream:
    rows = list(csv.DictReader(stream))
for row in rows:
    cells = table.add_row().cells
    values = [row["representation"], {"LogisticRegression": "Regresión logística", "RandomForest": "Random Forest", "XGBoost": "XGBoost"}[row["model"]]]
    for metric in ["accuracy", "f1", "sensitivity", "specificity"]:
        values.append(f"{100*float(row[metric+'_mean']):.2f} ± {100*float(row[metric+'_std']):.2f}".replace(".", ","))
    for cell, width, value in zip(cells, widths, values):
        cell.width = Inches(width)
        cell.text = value
for i, row in enumerate(table.rows):
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)
            for run in paragraph.runs:
                run.font.size = Pt(8)
                run.bold = i == 0
        props = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ["top", "left", "bottom", "right"]:
            border = OxmlElement("w:" + edge)
            for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                border.set(qn("w:" + key), value)
            borders.append(border)
        props.append(borders)
        if i == 0:
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "EEEEEE")
            props.append(shade)
columns(2)
p("La regresión logística se probó con dos intensidades de regularización, C igual a 0,1 y 1. En Random Forest se compararon 100 árboles con profundidad máxima de 12 y un mínimo de dos muestras por hoja, frente a 150 árboles sin límite de profundidad y una muestra mínima por hoja. Para XGBoost se compararon 100 árboles de profundidad 3 y tasa de aprendizaje de 0,1, frente a 150 árboles de profundidad 5 y tasa de 0,05. XGBoost utilizó una fracción de 0,8 de las muestras y de las características por árbol.")
p("Se aplicaron pesos de clase calculados a partir del entrenamiento. La transformación de los SMILES, el tratamiento de valores ausentes y la preparación de los descriptores también se ajustaron únicamente con los datos de entrenamiento. Se mantuvo un umbral de decisión de 0,5. Este procedimiento se repitió con 100 semillas distintas, lo que produjo 600 evaluaciones finales y permitió observar la variación entre particiones.")
p("La Tabla I y la Fig. 1 muestran que PaDEL alcanzó un F1 medio superior al de SMILES en los tres modelos. XGBoost mejoró su accuracy de 92,19% a 93,46% y su F1 de 95,04% a 95,88% al usar PaDEL. Random Forest con PaDEL obtuvo el mayor F1, con 96,05%, aunque la diferencia frente a XGBoost fue de solo 0,17 puntos porcentuales. Esta comparación describe los promedios observados; no demuestra una superioridad estadística.")

columns(1, new_page=True)
figure("reports/imagenes/matrices_todas_porcentaje.png", "Fig. 2. Matrices de confusión de las seis combinaciones. Las filas indican la clase real y las columnas, la predicción. Cada fila se normalizó dentro de cada repetición y después se promediaron los porcentajes. La diagonal muestra los aciertos de cada clase.", 7.15)
columns(2)
heading("IV. INTERPRETACIÓN DE LOS RESULTADOS")
p("La accuracy resume la proporción total de aciertos. El F1-score combina la capacidad de recuperar moléculas activas con la proporción de predicciones activas que son correctas. La sensibilidad expresa cuántas activas se detectan y la especificidad, cuántas inactivas se reconocen. Estas medidas se calcularon por repetición y luego se resumieron mediante su media y desviación estándar. Las barras de error no representan intervalos de confianza.")
p("La Fig. 2 permite distinguir comportamientos que no se aprecian al mirar únicamente el F1. XGBoost con PaDEL identificó, en promedio, el 96,67% de las moléculas activas, pero reconoció el 81,62% de las inactivas. El 18,38% restante de las inactivas se clasificó como activo. En una partición típica de 183 moléculas, estos porcentajes corresponden a promedios de 139,21 activas correctamente detectadas, 4,79 activas omitidas, 31,83 inactivas reconocidas y 7,17 inactivas clasificadas como activas.")
column_break()
p("En XGBoost, la especificidad media fue la misma con ambas representaciones, mientras que PaDEL redujo la proporción de activas omitidas de 4,94% a 3,33%. La mejora observada se concentró, por tanto, en recuperar mejor las moléculas activas. La regresión logística con PaDEL mostró otro equilibrio: alcanzó la mayor especificidad, 84,44%, con una sensibilidad de 93,73%. La elección del modelo depende de cuál de estos errores se considere más costoso en la etapa de selección.")
p("El conjunto contiene más moléculas activas que inactivas. Una regla que clasificara todas como activas obtendría una accuracy de 78,69% y un F1 de 88,07%, pese a tener especificidad nula. Los modelos evaluados superaron esa referencia y reconocieron una parte considerable de las inactivas. Aun así, la diferencia entre sensibilidad y especificidad confirma que una cifra global alta debe acompañarse de una lectura por clase.")

columns(2, new_page=True)
heading("V. DISCUSIÓN")
p("Los resultados sugieren que las características de las moléculas incluidas contienen información útil para distinguir los dos niveles de actividad definidos. PaDEL ofreció una mejora consistente en F1 frente a la representación textual utilizada. Esto no implica que cualquier uso de SMILES produzca un resultado inferior: aquí se evaluó una vectorización concreta de fragmentos de caracteres. Tampoco se examinó si una representación más compleja o una búsqueda más amplia de configuraciones cambiaría la comparación.")
p("La diferencia entre Random Forest y XGBoost con PaDEL fue pequeña respecto de la variación observada entre particiones. Ambos mostraron sensibilidad alta y una especificidad más moderada. Si el interés principal consiste en conservar posibles moléculas activas para revisarlas después, este comportamiento resulta pertinente. Cuando el objetivo es reducir el número de inactivas que avanzan a una etapa posterior, la mayor especificidad de la regresión logística merece atención, aunque se acompaña de más activas omitidas. Los resultados permiten reconocer esa decisión, pero no fijan su costo experimental.")
p("La evaluación se hizo con separaciones aleatorias. Aunque una misma estructura normalizada no aparece a la vez en entrenamiento y prueba dentro de una repetición, moléculas de familias químicas similares pueden quedar en ambos grupos. Por ello, el rendimiento puede ser más favorable que el que se obtendría al trabajar con familias completamente nuevas. Las 100 repeticiones reutilizan el mismo conjunto y se solapan; su dispersión refleja sensibilidad a la partición, no evidencia proveniente de 100 estudios independientes.")
p("La preparación de los datos también condiciona el alcance del resultado. Se usaron mediciones exactas y se retiró la zona intermedia entre los umbrales de actividad. Esta decisión permitió comparar clases bien delimitadas, pero deja sin evaluar compuestos próximos a esos límites. Además, las mediciones proceden de diferentes ensayos y pueden conservar variaciones de sus condiciones originales. La mediana resume registros repetidos, pero no elimina toda esa heterogeneidad.")
p("El siguiente paso razonable sería evaluar los modelos con una separación por familias estructurales o con un conjunto externo. También convendría revisar el umbral de decisión según el propósito de selección y ampliar de forma controlada la exploración de configuraciones. Estas acciones permitirían determinar si el comportamiento se mantiene fuera de las condiciones evaluadas. La confirmación de actividad y la valoración de seguridad requieren evidencia experimental adicional.")
column_break()
heading("VI. CONCLUSIONES")
p("Se construyó y evaluó un flujo de clasificación de bioactividad frente a MMP-9 humana con 913 moléculas, dos representaciones y tres modelos. XGBoost con PaDEL alcanzó 93,46% de accuracy y 95,88% de F1, con una sensibilidad de 96,67% y una especificidad de 81,62%. Su desempeño fue cercano al de Random Forest con PaDEL, que presentó el mayor F1 medio del experimento.")
p("La principal conclusión es que el rendimiento debe interpretarse por clase. Los modelos recuperaron una proporción alta de activas, pero mantuvieron errores entre las inactivas. La comparación ofrece una base reproducible para priorización computacional dentro del conjunto estudiado. Su aplicación a nuevas moléculas exige comprobar primero la generalización y conservar la distinción entre una predicción de actividad y una validación farmacológica.")
heading("VII. TRAZABILIDAD DEL ANÁLISIS")
p("Los resultados se respaldan en los archivos del proyecto [5]. La descarga y curación se encuentran en mmp9/data.py, el cálculo PaDEL en mmp9/features.py y el entrenamiento y las métricas en mmp9/experiment.py. Las figuras proceden de mmp9/plots.py. La auditoría de mmp9/verify.py verificó las 600 evaluaciones y 109.800 filas de predicción, incluyendo la separación de datos y la reproducción de resultados al cargar los modelos. Estas filas incluyen apariciones repetidas de moléculas y no corresponden a compuestos distintos.")
heading("REFERENCIAS")
references = [
    '[1] UniProt Consortium, “Matrix metalloproteinase-9, Homo sapiens, P14780,” UniProtKB. [En línea]. Disponible: https://www.uniprot.org/uniprotkb/P14780/entry. Consulta: 22 sep. 2026.',
    '[2] EMBL-EBI, “ChEMBL, versión 37. Objetivo CHEMBL321 y registros de bioactividad IC50,” 2026. [En línea]. Disponible: https://www.ebi.ac.uk/chembl/. Datos descargados el 22 sep. 2026, licencia CC BY-SA 3.0.',
    '[3] C. W. Yap, “PaDEL-descriptor: An open source software to calculate molecular descriptors and fingerprints,” Journal of Computational Chemistry, vol. 32, no. 7, pp. 1466–1474, 2011, doi: 10.1002/jcc.21707.',
    '[4] T. Chen y C. Guestrin, “XGBoost: A Scalable Tree Boosting System,” en Proc. 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, 2016, pp. 785–794, doi: 10.1145/2939672.2939785.',
    '[5] G. Ortiz Porras y S. T. Rubiano Meneses, “Proyecto de clasificación de bioactividad frente a MMP-9,” código Python, datos curados y resultados del taller, Universidad del Rosario, 2026. Archivos locales: mmp9/, results/main/ y reports/imagenes/.',
]
for text in [references[i] for i in [0, 3, 1, 2, 4]]:
    paragraph = p(text)
    paragraph.paragraph_format.first_line_indent = Inches(-.18)
    paragraph.paragraph_format.left_indent = Inches(.18)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        run.font.size = Pt(8)

doc.core_properties.title = "Clasificación de la actividad de moléculas frente a MMP-9 mediante aprendizaje automático"
doc.core_properties.author = "Geronimo Ortiz Porras; Samuel Tomas Rubiano Meneses"
doc.core_properties.subject = "Artículo de conferencia IEEE sobre bioactividad de MMP-9"
doc.core_properties.keywords = "MMP-9, XGBoost, PaDEL, SMILES, bioactividad"
doc.core_properties.comments = ""
doc.save(OUT / "informe_IEEE_MMP9.docx")
print(OUT / "informe_IEEE_MMP9.docx")
